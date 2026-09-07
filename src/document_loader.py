from __future__ import annotations
from dataclasses import dataclass, asdict
from pathlib import Path
from typing import List, Dict
import hashlib
import re

@dataclass
class DocumentPage:
    text: str
    source: str
    page: int
    file_type: str
    section: str = ""
    document_id: str = ""

def _clean(text: str) -> str:
    text = text.replace("\x00", " ")
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()

def _doc_id(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()[:16]

def load_pdf(path: Path) -> List[DocumentPage]:
    from pypdf import PdfReader
    reader = PdfReader(str(path))
    did = _doc_id(path)
    pages = []
    for i, page in enumerate(reader.pages, start=1):
        text = _clean(page.extract_text() or "")
        if text:
            pages.append(DocumentPage(text, path.name, i, "pdf", document_id=did))
    return pages

def load_docx(path: Path) -> List[DocumentPage]:
    from docx import Document
    doc = Document(str(path))
    did = _doc_id(path)
    blocks = []
    for p in doc.paragraphs:
        if p.text.strip():
            blocks.append(p.text.strip())
    for table in doc.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells]
            blocks.append(" | ".join(cells))
    text = _clean("\n".join(blocks))
    return [DocumentPage(text, path.name, 1, "docx", document_id=did)] if text else []

def load_txt(path: Path) -> List[DocumentPage]:
    text = _clean(path.read_text(encoding="utf-8", errors="ignore"))
    did = _doc_id(path)
    return [DocumentPage(text, path.name, 1, "txt", document_id=did)] if text else []

def load_documents(directory: Path) -> List[DocumentPage]:
    pages: List[DocumentPage] = []
    for path in sorted(directory.rglob("*")):
        if not path.is_file():
            continue
        suffix = path.suffix.lower()
        try:
            if suffix == ".pdf":
                pages.extend(load_pdf(path))
            elif suffix == ".docx":
                pages.extend(load_docx(path))
            elif suffix in {".txt", ".md"}:
                pages.extend(load_txt(path))
        except Exception as exc:
            print(f"[WARN] Failed to load {path}: {exc}")
    return pages
